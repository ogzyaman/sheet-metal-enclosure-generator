"""Copied from cadkit/manifest.py -- this repo does not depend on cadkit.

Manifest CSV/XLSX writer. Deliberate duplication, same pattern used
across this author's portfolio repos: public/standalone repos never
link against the private lib/cadkit package.
"""
import csv

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NUMERIC_FIELDS = ["length_mm", "width_mm", "height_mm", "hole_dia_mm", "fillet_r_mm"]

HEADER_FONT = Font(name="Calibri", size=11, bold=True)
HEADER_FILL = PatternFill(start_color="FFD9D9D9", end_color="FFD9D9D9", fill_type="solid")
DATA_FONT = Font(name="Calibri", size=11)
STATUS_OK_FONT = Font(name="Calibri", size=11, color="FF008000")
STATUS_FAILED_FONT = Font(name="Calibri", size=11, color="FFFF0000")
THIN_SIDE = Side(style="thin", color="FF000000")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def build_manifest_fields(formats, numeric_fields=NUMERIC_FIELDS):
    fields = ["variant_id"] + list(numeric_fields)
    for fmt in formats:
        fields.append(f"{fmt}_file")
    fields.append("status")
    return fields


def _to_number(value):
    """Returns int for whole numbers, float otherwise -- so Excel shows 60, not 60.0."""
    f = float(value)
    return int(f) if f.is_integer() else f


def write_manifest_csv(manifest_rows, manifest_path, manifest_fields):
    with open(manifest_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=manifest_fields)
        writer.writeheader()
        writer.writerows(manifest_rows)


def write_manifest_xlsx(manifest_rows, xlsx_path, manifest_fields, numeric_fields=NUMERIC_FIELDS):
    wb = Workbook()
    ws = wb.active
    ws.title = "manifest"

    numeric_fields = set(numeric_fields)

    ws.append(manifest_fields)
    for col_idx, field in enumerate(manifest_fields, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="right" if field in numeric_fields else "left")

    for row in manifest_rows:
        values = [_to_number(row[f]) if f in numeric_fields else row[f] for f in manifest_fields]
        ws.append(values)

    for row_offset, row in enumerate(manifest_rows):
        row_idx = row_offset + 2
        for col_idx, field in enumerate(manifest_fields, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="right" if field in numeric_fields else "left")
            if field == "status":
                if row["status"] == "ok":
                    cell.font = STATUS_OK_FONT
                elif row["status"].startswith("failed"):
                    cell.font = STATUS_FAILED_FONT
                else:
                    cell.font = DATA_FONT
            else:
                cell.font = DATA_FONT

    ws.freeze_panes = "A2"

    for col_idx, field in enumerate(manifest_fields, start=1):
        cell_lengths = [len(str(field))] + [len(str(row[field])) for row in manifest_rows]
        column_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[column_letter].width = max(cell_lengths) + 2

    wb.save(xlsx_path)
